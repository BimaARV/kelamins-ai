"""DOCX reading and writing via python-docx."""

from __future__ import annotations

import logging
import os

from app.documents.extractor import DocumentResult, normalize_text

logger = logging.getLogger(__name__)


def read_docx(file_path: str) -> DocumentResult:
    """Extract text, paragraphs and headings from a .docx file."""
    from docx import Document as DocxDocument

    doc = DocxDocument(file_path)
    core = doc.core_properties

    metadata: dict = {
        "title": core.title or "",
        "author": core.author or "",
        "paragraphs": len(doc.paragraphs),
        "file_path": file_path,
    }

    paragraphs_text: list[str] = []
    sections: list[dict] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        paragraphs_text.append(text)
        if para.style and para.style.name.startswith("Heading"):
            sections.append({"heading": para.style.name, "text": text[:200]})

    text = normalize_text("\n\n".join(paragraphs_text))
    return DocumentResult(text=text, metadata=metadata, sections=sections)


def write_docx_report(report_data: dict, output_path: str) -> str:
    """Generate a .docx report from structured data."""
    from docx import Document as DocxDocument
    from docx.shared import Pt, Inches

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

    doc = DocxDocument()
    style = doc.styles["Normal"]
    style.font.size = Pt(10)

    doc.add_heading(report_data.get("title", "KELA AI Report"), level=0)

    gen_at = report_data.get("generated_at", "")
    if gen_at:
        doc.add_paragraph(f"Generated: {gen_at}").style = doc.styles["Normal"]

    summary = report_data.get("executive_summary", "")
    if summary:
        doc.add_heading("Executive Summary", level=1)
        doc.add_paragraph(summary)

    events = report_data.get("events", [])
    if events:
        doc.add_heading("Events", level=1)
        table = doc.add_table(rows=1, cols=5, style="Table Grid")
        hdr = table.rows[0].cells
        for i, label in enumerate(["#", "Type", "Title", "Confidence", "Sources"]):
            hdr[i].text = label
        for idx, ev in enumerate(events, 1):
            row = table.add_row().cells
            row[0].text = str(idx)
            row[1].text = ev.get("event_type", "")
            row[2].text = ev.get("title", "")[:60]
            row[3].text = ev.get("confidence", "")
            row[4].text = str(ev.get("article_count", ""))

    appendix = report_data.get("appendix")
    if appendix:
        doc.add_heading("Appendix", level=1)
        for key, value in appendix.items():
            doc.add_paragraph(f"{key}: {value}")

    doc.save(output_path)
    return output_path
